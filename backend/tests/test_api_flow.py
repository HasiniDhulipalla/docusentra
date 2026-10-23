"""End-to-end authorization flow smoke test using an isolated SQLite database."""
import os
from io import BytesIO
import shutil
from docx import Document as WordDocument
from pathlib import Path

_test_id = str(os.getpid())
_test_db = Path(f"/tmp/trustvault-integration-{_test_id}.db")
_test_storage = Path(f"/tmp/trustvault-integration-storage-{_test_id}")


def _minimal_pdf(text="TrustVault secure preview test"):
    """Build a syntactically valid one-page PDF fixture without extra test dependencies."""
    safe_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 18 Tf 24 82 Td ({safe_text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 360 144] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode("ascii"))
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii"))
    return bytes(pdf)

_test_db.unlink(missing_ok=True)
os.environ["DATABASE_URL"] = f"sqlite:////tmp/trustvault-integration-{_test_id}.db"
os.environ["STORAGE_DIR"] = str(_test_storage)
os.environ["TOKEN_SECRET"] = "test-only-secret-for-integration-tests"

from fastapi.testclient import TestClient
import app.main as api_module
from app.main import app


def test_share_download_revoke_security_flow():
    with TestClient(app) as client:
        # Protected resources require authentication; weak registration passwords fail closed.
        assert client.get("/api/documents").status_code == 401
        weak_password = client.post("/api/auth/register", json={
            "username": "weak_test", "email": "weak@test.example", "password": "Short123!",
        })
        assert weak_password.status_code == 422
        malformed_email = client.post("/api/auth/register", json={
            "username": "badmail_test", "email": "not-an-email", "password": "Strong-Enough-123!",
        })
        assert malformed_email.status_code == 422
        alice_response = client.post("/api/auth/register", json={
            "username": "alice_test", "email": "alice@test.example", "password": "Alice-Test-123",
        })
        assert alice_response.status_code == 201, alice_response.text
        alice_token = alice_response.json()["access_token"]
        alice_headers = {"Authorization": f"Bearer {alice_token}"}

        bob_response = client.post("/api/auth/register", json={
            "username": "bob_test", "email": "bob@test.example", "password": "Bob-Test-1234",
        })
        assert bob_response.status_code == 201, bob_response.text
        bob_headers = {"Authorization": f"Bearer {bob_response.json()['access_token']}"}
        assert client.get("/api/users?q=b", headers=alice_headers).json() == []
        exact_user_lookup = client.get("/api/users?q=BOB_TEST", headers=alice_headers)
        assert exact_user_lookup.status_code == 200
        assert [item["username"] for item in exact_user_lookup.json()] == ["bob_test"]
        duplicate_account = client.post("/api/auth/register", json={
            "username": "ALICE_TEST", "email": "another@test.example", "password": "Another-Strong-123!",
        })
        assert duplicate_account.status_code == 409

        # Extension allowlisting alone is insufficient: mismatched content must be rejected.
        spoofed_pdf = client.post("/api/documents", headers=alice_headers,
                                  files={"file": ("pretend.pdf", b"MZ-not-a-pdf", "application/pdf")})
        assert spoofed_pdf.status_code == 415
        forbidden_extension = client.post("/api/documents", headers=alice_headers,
                                          files={"file": ("payload.exe", b"MZ", "application/octet-stream")})
        assert forbidden_extension.status_code == 415
        invalid_category = client.post("/api/documents", headers=alice_headers,
                                       files={"file": ("category.txt", b"category validation test", "text/plain")},
                                       data={"category": "Payroll"})
        assert invalid_category.status_code == 400
        original_max_upload = api_module.MAX_UPLOAD_BYTES
        try:
            api_module.MAX_UPLOAD_BYTES = 16
            oversized = client.post("/api/documents", headers=alice_headers,
                                    files={"file": ("too-large.txt", b"This upload is over sixteen bytes.", "text/plain")})
            assert oversized.status_code == 413
        finally:
            api_module.MAX_UPLOAD_BYTES = original_max_upload

        word_buffer = BytesIO()
        word_doc = WordDocument()
        word_doc.add_heading("Confidential Project Plan", level=1)
        word_doc.add_paragraph("Reader users must be able to read this text but cannot download the original Word file.")
        word_doc.save(word_buffer)
        docx_bytes = word_buffer.getvalue()
        upload = client.post("/api/documents", headers=alice_headers,
                             files={"file": ("private.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert upload.status_code == 201, upload.text
        document_id = upload.json()["id"]

        # A valid login is not sufficient: the document is private by default.
        denied = client.get(f"/api/documents/{document_id}/download", headers=bob_headers)
        assert denied.status_code == 403
        denied_preview = client.get(f"/api/documents/{document_id}/preview", headers=bob_headers)
        assert denied_preview.status_code == 403

        share = client.post(f"/api/documents/{document_id}/shares", headers=alice_headers,
                            json={"user_identifier": "bob_test", "role": "reader"})
        assert share.status_code == 201, share.text
        assert client.get("/api/documents", headers=bob_headers).json()[0]["access_role"] == "reader"
        # Unknown/corrupt persisted roles must never fall through to Editor-like permissions.
        test_session = api_module.SessionLocal()
        try:
            share_row = test_session.get(api_module.DocumentShare, share.json()["id"])
            share_row.role = "administrator"
            test_session.commit()
        finally:
            test_session.close()
        corrupt_role_download = client.get(f"/api/documents/{document_id}/download", headers=bob_headers)
        assert corrupt_role_download.status_code == 403
        test_session = api_module.SessionLocal()
        try:
            share_row = test_session.get(api_module.DocumentShare, share.json()["id"])
            share_row.role = "reader"
            test_session.commit()
        finally:
            test_session.close()
        reader_rename = client.patch(f"/api/documents/{document_id}", headers=bob_headers,
                                     json={"filename": "renamed.docx"})
        assert reader_rename.status_code == 403
        reader_share = client.post(f"/api/documents/{document_id}/shares", headers=bob_headers,
                                   json={"user_identifier": "alice_test", "role": "reader"})
        assert reader_share.status_code == 403
        reader_delete = client.delete(f"/api/documents/{document_id}", headers=bob_headers)
        assert reader_delete.status_code == 403
        reader_shares_list = client.get(f"/api/documents/{document_id}/shares", headers=bob_headers)
        assert reader_shares_list.status_code == 403
        reader_preview = client.get(f"/api/documents/{document_id}/preview", headers=bob_headers)
        assert reader_preview.status_code == 200, reader_preview.text
        preview_data = reader_preview.json()
        assert preview_data["preview_type"] in {"blocks", "pages"}
        if preview_data["preview_type"] == "blocks":
            assert any("Confidential Project Plan" in block.get("text", "") for block in preview_data["blocks"])
            assert any("cannot download the original Word file" in block.get("text", "") for block in preview_data["blocks"])
        else:
            assert preview_data["page_count"] >= 1
            converted_preview = client.get(f"/api/documents/{document_id}/preview-content", headers=bob_headers)
            assert converted_preview.status_code == 415
            first_word_page = client.get(f"/api/documents/{document_id}/preview-pages/1", headers=bob_headers)
            assert first_word_page.status_code == 200, first_word_page.text
            assert first_word_page.content.startswith(b"\x89PNG\r\n\x1a\n")
            assert not first_word_page.content.startswith(b"%PDF")

        # PDF/Office pages are delivered as unwatermarked raster images rather than raw PDF bytes.
        # Image preview bytes are safely re-encoded without adding visible watermarking.
        sample_pdf = _minimal_pdf()
        sample_png = bytes.fromhex("89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000b49444154789c636000020000050001a5f645400000000049454e44ae426082")
        for filename, payload, mime, expected_preview_type in (
            ("sample.pdf", sample_pdf, "application/pdf", "pages"),
            ("sample.png", sample_png, "image/png", "image"),
        ):
            additional_upload = client.post("/api/documents", headers=alice_headers,
                files={"file": (filename, payload, mime)})
            assert additional_upload.status_code == 201, additional_upload.text
            extra_id = additional_upload.json()["id"]
            extra_share = client.post(f"/api/documents/{extra_id}/shares", headers=alice_headers,
                json={"user_identifier": "bob_test", "role": "reader"})
            assert extra_share.status_code == 201, extra_share.text
            extra_preview = client.get(f"/api/documents/{extra_id}/preview", headers=bob_headers)
            assert extra_preview.status_code == 200, extra_preview.text
            assert extra_preview.json()["preview_type"] == expected_preview_type
            if expected_preview_type == "pages":
                assert extra_preview.json()["page_count"] == 1
                preview_content = client.get(f"/api/documents/{extra_id}/preview-content", headers=bob_headers)
                assert preview_content.status_code == 415
                preview_content = client.get(f"/api/documents/{extra_id}/preview-pages/1", headers=bob_headers)
                assert preview_content.status_code == 200, preview_content.text
                assert preview_content.content.startswith(b"\x89PNG\r\n\x1a\n")
                assert not preview_content.content.startswith(b"%PDF")
                assert client.get(f"/api/documents/{extra_id}/preview-pages/2", headers=bob_headers).status_code == 404
            else:
                preview_content = client.get(f"/api/documents/{extra_id}/preview-content", headers=bob_headers)
                assert preview_content.status_code == 200, preview_content.text
                assert preview_content.content.startswith(b"\x89PNG\r\n\x1a\n")
                assert preview_content.content != payload
                if expected_preview_type == "image":
                    # Decoding and re-encoding strips metadata but preserves visible image pixels.
                    with api_module.Image.open(BytesIO(payload)) as source_image:
                        source_pixel = source_image.convert("RGB").getpixel((0, 0))
                    with api_module.Image.open(BytesIO(preview_content.content)) as preview_image:
                        assert preview_image.convert("RGB").getpixel((0, 0)) == source_pixel
            assert preview_content.headers["content-type"].startswith("image/png")
            assert "no-store" in preview_content.headers.get("cache-control", "")
            assert preview_content.headers.get("x-content-type-options") == "nosniff"
            assert preview_content.headers.get("x-frame-options") == "SAMEORIGIN"
            assert client.get(f"/api/documents/{extra_id}/download", headers=bob_headers).status_code == 403

        # Reader permission allows preview endpoints but never original-file download.
        blocked_reader_download = client.get(f"/api/documents/{document_id}/download", headers=bob_headers)
        assert blocked_reader_download.status_code == 403

        # The owner upgrades the permission; editors may rename metadata.
        upgrade = client.post(f"/api/documents/{document_id}/shares", headers=alice_headers,
                              json={"user_identifier": "bob_test", "role": "editor"})
        assert upgrade.status_code == 201
        editor_rename = client.patch(f"/api/documents/{document_id}", headers=bob_headers,
                                     json={"filename": "renamed.docx"})
        assert editor_rename.status_code == 200, editor_rename.text

        allowed = client.get(f"/api/documents/{document_id}/download", headers=bob_headers)
        assert allowed.status_code == 200
        assert allowed.content == docx_bytes
        editor_cannot_grant = client.post(f"/api/documents/{document_id}/shares", headers=bob_headers,
                                          json={"user_identifier": "alice_test", "role": "reader"})
        assert editor_cannot_grant.status_code == 403
        editor_cannot_delete = client.delete(f"/api/documents/{document_id}", headers=bob_headers)
        assert editor_cannot_delete.status_code == 403

        revoke = client.delete(f"/api/documents/{document_id}/shares/{share.json()['id']}", headers=alice_headers)
        assert revoke.status_code == 204
        blocked_after_revoke = client.get(f"/api/documents/{document_id}/download", headers=bob_headers)
        assert blocked_after_revoke.status_code == 403
        blocked_preview_after_revoke = client.get(f"/api/documents/{document_id}/preview-content", headers=bob_headers)
        assert blocked_preview_after_revoke.status_code == 403
        assert client.get(f"/api/documents/{document_id}/preview-pages/1", headers=bob_headers).status_code == 403

        owner_activity = client.get("/api/activity", headers=alice_headers)
        assert owner_activity.status_code == 200
        actions = [event["action"] for event in owner_activity.json()]
        assert "ACCESS_GRANTED" in actions
        assert "ACCESS_REVOKED" in actions
        assert actions.count("ACCESS_DENIED") >= 2
        assert "DOCUMENT_PREVIEWED" in actions

        # An upload is private by default and the owner can delete it. Keep storage cleanup explicit.
        text_upload = client.post("/api/documents", headers=alice_headers,
                                  files={"file": ("notes.txt", b"Private test notes\n", "text/plain")})
        assert text_upload.status_code == 201, text_upload.text
        text_id = text_upload.json()["id"]
        test_session = api_module.SessionLocal()
        try:
            text_doc = test_session.get(api_module.Document, text_id)
            assert text_doc is not None
            storage_key = text_doc.storage_key
        finally:
            test_session.close()
        storage_file = _test_storage / storage_key
        assert storage_file.is_file()
        text_preview = client.get(f"/api/documents/{text_id}/preview", headers=alice_headers)
        assert text_preview.status_code == 200, text_preview.text
        assert text_preview.json()["preview_type"] == "pages"
        assert text_preview.json().get("blocks", []) == []
        text_preview_page = client.get(f"/api/documents/{text_id}/preview-pages/1", headers=alice_headers)
        assert text_preview_page.status_code == 200, text_preview_page.text
        assert text_preview_page.content.startswith(b"\x89PNG\r\n\x1a\n")
        cache_dir = _test_storage / ".preview_cache"
        cache_dir.mkdir(parents=True, exist_ok=True)
        real_text_cache = list(cache_dir.glob(f"{storage_key}-*-text-page-v2-1.png"))
        assert real_text_cache, "text preview should use a private cached raster page"
        fake_cached_preview = cache_dir / f"{storage_key}-17-23.pdf"
        fake_cached_page = cache_dir / f"{storage_key}-17-23-page-1.png"
        fake_cached_preview.write_bytes(b"%PDF-private-render")
        fake_cached_page.write_bytes(b"private-png-render")
        delete_response = client.delete(f"/api/documents/{text_id}", headers=alice_headers)
        assert delete_response.status_code == 204
        assert not storage_file.exists()
        assert not fake_cached_preview.exists()
        assert not fake_cached_page.exists()

        # Ten login attempts consume the sliding window; the next attempt receives 429 before password verification.
        api_module.clear_rate_limit("login-attempt", "testclient")
        for _ in range(api_module.LOGIN_MAX_FAILURES):
            bad_login = client.post("/api/auth/login", json={"identifier": "alice_test", "password": "Wrong-password-123"})
            assert bad_login.status_code == 401
        limited_login = client.post("/api/auth/login", json={"identifier": "alice_test", "password": "Wrong-password-123"})
        assert limited_login.status_code == 429
        assert "Retry-After" in limited_login.headers
        # A locked-out IP is rejected before password verification; after cooldown/admin reset the correct password works.
        locked_valid_login = client.post("/api/auth/login", json={"identifier": "alice_test", "password": "Alice-Test-123"})
        assert locked_valid_login.status_code == 429
        api_module.clear_rate_limit("login-attempt", "testclient")
        valid_login = client.post("/api/auth/login", json={"identifier": "alice_test", "password": "Alice-Test-123"})
        assert valid_login.status_code == 200
        owner_activity_after_auth_tests = client.get("/api/activity", headers=alice_headers)
        assert owner_activity_after_auth_tests.status_code == 200
        assert "LOGIN_FAILED" in [event["action"] for event in owner_activity_after_auth_tests.json()]

        # Password changes require the correct current password and a strong, distinct replacement.
        bad_current_password = client.post("/api/auth/change-password", headers=alice_headers, json={
            "current_password": "Wrong-current-password", "new_password": "Alice-New-Pass-1234",
        })
        assert bad_current_password.status_code == 400
        weak_new_password = client.post("/api/auth/change-password", headers=alice_headers, json={
            "current_password": "Alice-Test-123", "new_password": "short",
        })
        assert weak_new_password.status_code == 422
        changed_password = client.post("/api/auth/change-password", headers=alice_headers, json={
            "current_password": "Alice-Test-123", "new_password": "Alice-New-Pass-1234",
        })
        assert changed_password.status_code == 200, changed_password.text
        assert client.post("/api/auth/login", json={"identifier": "alice_test", "password": "Alice-Test-123"}).status_code == 401
        api_module.clear_rate_limit("login-attempt", "testclient")
        assert client.post("/api/auth/login", json={"identifier": "alice_test", "password": "Alice-New-Pass-1234"}).status_code == 200

    # Keep the isolated database file until the full pytest process exits; later challenge
    # tests share this imported SQLAlchemy engine and need its SQLite path to remain valid.
